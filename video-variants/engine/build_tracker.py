"""Build the BANK tracker (Google Sheet), bank_index.json and README_FOR_AI.md."""
from __future__ import annotations

import datetime as dt
import json
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TRACKER_NAME = "Katie Video Bank Tracker"
STATUSES = ["Not posted", "Posted", "Skip", "Needs fix"]
GROUP_LABEL = {"FRESH": "Fresh Veo 3", "REUSED": "Reused Veo 3", "SEED": "Seedance 2.5"}


def file_url(i):
    return f"https://drive.google.com/file/d/{i}/view" if i else ""


def folder_url(i):
    return f"https://drive.google.com/drive/folders/{i}" if i else ""


def one_line(h: dict) -> str:
    body = " ".join(h["body"].split())
    return " / ".join(x for x in (h["headline"], body, h["cta"]) if x)


def collect(sources, rc, remote, bank_id, work):
    files = json.loads(rc("lsjson", "-R", "--files-only", remote(bank_id)))
    dirs = json.loads(rc("lsjson", "-R", "--dirs-only", remote(bank_id)))
    fid = {f["Path"]: f["ID"] for f in files}
    did = {d["Path"]: d["ID"] for d in dirs}
    rows = []
    for s in sources:
        folder = f"{s['source_id']}_{s['slug']}"
        jpath = os.path.join(work, "bank", s["bank_folder"], folder, f"{folder}.json")
        if not os.path.exists(jpath):
            continue
        res = json.load(open(jpath))
        base = f"{s['bank_folder']}/{folder}"
        entry = {
            "source_id": s["source_id"], "group": s["group"], "group_label": GROUP_LABEL[s["group"]],
            "theme": s["theme"], "note": s["note"], "original_name": s["original_name"],
            "original_drive_id": s["original_drive_id"], "hook": s["hook"],
            "folder": base, "folder_drive_id": did.get(base),
            "preview_drive_id": fid.get(f"_previews/{folder}_preview.jpg"),
            "data_drive_id": fid.get(f"_data/{folder}.json"),
            "variants": [],
        }
        for v in res["variants"]:
            path = f"{base}/{v['file']}"
            entry["variants"].append({
                "video_id": v["video_id"], "file": v["file"], "path": path, "drive_id": fid.get(path),
                "url": file_url(fid.get(path)), "variant": v["variant"], "style": v["style"],
                "family": v["family"], "size_pct": v["size_pct"], "headline_color": v["headline_color"],
                "body_color": v["body_color"], "cta_box": v["cta_box"], "text_box_xywh": v["text_box"],
                "duration_s": v["duration_s"], "bytes": v["bytes"],
            })
        rows.append(entry)
    return rows, fid, did


def build_xlsx(rows, path):
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="222222")

    ws = wb.active
    ws.title = "Videos"
    cols = [("Video ID", 18), ("Status", 13), ("Posted to (account)", 22), ("Date posted", 13),
            ("Source ID", 13), ("Variant", 8), ("Group", 14), ("Theme", 10), ("On-screen text", 60),
            ("Text style", 42), ("Video", 10), ("File name", 42), ("Folder", 44), ("Drive file ID", 36),
            ("Length (s)", 10), ("Notes", 36)]
    ws.append([c for c, _ in cols])
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.cell(1, i).font, ws.cell(1, i).fill = head_font, head_fill
    n = 1
    for e in rows:
        for v in e["variants"]:
            n += 1
            ws.append([v["video_id"], "Not posted", "", "", e["source_id"], v["variant"], e["group_label"],
                       e["theme"], one_line(e["hook"]), v["style"],
                       f'=HYPERLINK("{v["url"]}","Open")' if v["url"] else "", v["file"], e["folder"],
                       v["drive_id"] or "", round(v["duration_s"] or 0, 1),
                       e["note"] if e["group"] == "REUSED" else ""])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:P{n}"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{max(n, 2)}")
    rng = f"A2:P{max(n, 2)}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$B2="Posted"'], fill=PatternFill("solid", fgColor="C8E6C9")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$B2="Skip"'], fill=PatternFill("solid", fgColor="E0E0E0")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$B2="Needs fix"'], fill=PatternFill("solid", fgColor="FFCDD2")))

    src = wb.create_sheet("Clips")
    scols = [("Source ID", 13), ("Group", 14), ("Theme", 10), ("Original clip", 36), ("Posted", 8),
             ("Variants", 9), ("On-screen text", 60), ("Preview (all 10)", 16), ("Folder", 12), ("Notes", 36)]
    src.append([c for c, _ in scols])
    for i, (_, w) in enumerate(scols, 1):
        src.column_dimensions[src.cell(1, i).column_letter].width = w
        src.cell(1, i).font, src.cell(1, i).fill = head_font, head_fill
    for r, e in enumerate(rows, 2):
        src.append([e["source_id"], e["group_label"], e["theme"], e["original_name"],
                    f'=COUNTIFS(Videos!$E:$E,A{r},Videos!$B:$B,"Posted")', len(e["variants"]),
                    one_line(e["hook"]),
                    f'=HYPERLINK("{file_url(e["preview_drive_id"])}","See preview")' if e["preview_drive_id"] else "",
                    f'=HYPERLINK("{folder_url(e["folder_drive_id"])}","Open")' if e["folder_drive_id"] else "",
                    e["note"]])
    src.freeze_panes = "B2"

    how = wb.create_sheet("How to use")
    how.column_dimensions["A"].width = 120
    for line in HOW_TO:
        how.append([line])
    for row in how.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    how["A1"].font = Font(bold=True, size=14)
    wb.save(path)


HOW_TO = [
    "Katie Video Bank - how to use this sheet",
    "Every row on the Videos tab is one finished video. Change Status to Posted after you post it "
    "(the row turns green) and fill in the account and date. Use Skip for ones you don't want, Needs fix for problems.",
    "Filter the Status column to 'Not posted' to see what's left. Sort or filter by Group, Theme or Source ID.",
    "Each clip has 10 variants (v01-v10): same video, different text font/size/colour and an invisible colour tweak. "
    "Spread variants of the same clip across different days/accounts.",
    "The Clips tab shows one row per original clip, how many of its variants are posted, and a preview image of all 10.",
    "Folders in BANK: 1_FRESH_veo3, 2_REUSED_veo3 (clips posted before), 3_SEEDANCE_2.5; one sub-folder per clip. "
    "_previews has a contact sheet per clip. _data and bank_index.json are for AI tools.",
    "AI tools: read README_FOR_AI.md in the BANK folder first.",
]


def readme(rows, tracker_id, bank_id, generated):
    total = sum(len(e["variants"]) for e in rows)
    by = {}
    for e in rows:
        by.setdefault(e["group_label"], [0, 0])
        by[e["group_label"]][0] += 1
        by[e["group_label"]][1] += len(e["variants"])
    groups = "\n".join(f"- {g}: {c} clip{'s' * (c != 1)} -> {v} videos" for g, (c, v) in by.items())
    return f"""# Katie Video Bank - README for AI tools

Generated {generated}. {total} videos from {len(rows)} source clip{'s' * (len(rows) != 1)}.
{groups}

## Where things are
- BANK folder id: `{bank_id}` (AI Influencers / katie / BANK)
- Tracker (Google Sheet): "{TRACKER_NAME}", id `{tracker_id or 'see BANK folder'}`
  - Tab **Videos**: one row per video. Column B **Status** is the source of truth:
    `Not posted` | `Posted` | `Skip` | `Needs fix`. Columns C/D = account posted to, date posted.
  - Tab **Clips**: one row per source clip with a preview link.
- `bank_index.json`: every clip and video with Drive file ids, on-screen text, style and folder.
  (It does not track posting; the sheet does.)
- `_previews/<clip>_preview.jpg`: all 10 variants of a clip side by side.
- `_data/<clip>.json`: full render details for a clip (text position, colour grade, checks).
- `_posted_log/`: fallback posting log (see below).

## Naming
- Video id: `<GROUP>-<CODE>-vNN`, e.g. `FRESH-R01-v03`, `REUSED-07-v10`, `SEED-02-v01`.
- File: `<GROUP>-<CODE>_<slug>_vNN.mp4` inside `<group folder>/<GROUP>-<CODE>_<slug>/`.
- Groups: FRESH = new Veo 3 clips; REUSED = Veo 3 clips that were posted before
  (their source was used already, so prefer FRESH/SEED when choosing); SEED = Seedance 2.5.
- All videos are 1080x1920 H.264, same length and audio as the source clip.

## Picking videos to post
1. Read the sheet (Videos tab). Only pick rows with Status `Not posted`.
2. Treat a video as posted if the sheet says `Posted` OR a matching file exists in `_posted_log/`.
3. Avoid posting two variants of the same Source ID to the same account within 7 days,
   and mix themes across a day's posts.
4. Get the file by its Drive file id (column N or `bank_index.json`).

## Marking a video as posted
- Preferred: edit the sheet row - Status `Posted`, account in column C, date (YYYY-MM-DD) in column D.
- If you cannot edit sheets, create an empty text file in `_posted_log/` named
  `<video_id>__<account>__<YYYY-MM-DD>.txt` (e.g. `FRESH-R01-v03__katieswellness__2026-09-25.txt`).
  Whoever next has sheet access should copy these into the sheet.
"""


def build_and_upload(sources, rc, remote, bank_id, work, log):
    rows, fid, _ = collect(sources, rc, remote, bank_id, work)
    generated = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    out = os.path.join(work, "tracker")
    os.makedirs(out, exist_ok=True)
    xlsx = os.path.join(out, f"{TRACKER_NAME}.xlsx")
    build_xlsx(rows, xlsx)

    top = json.loads(rc("lsjson", remote(bank_id)))
    existing = [f for f in top if f["Name"].startswith(TRACKER_NAME)]
    name = TRACKER_NAME if not existing else f"{TRACKER_NAME} - rebuilt {generated[:10]}"
    if existing:
        log(f"A tracker already exists (keeping it, since it holds posting marks); new one saved as '{name}'")
    rc("copyto", xlsx, remote(bank_id, f"{name}.xlsx"), "--drive-import-formats", "xlsx")
    top = json.loads(rc("lsjson", remote(bank_id)))
    tracker = next((f for f in top if f["Name"].startswith(name)), None)
    tracker_id = tracker["ID"] if tracker else None

    index = {
        "about": "Katie video bank: every finished video with Drive ids. Posting status lives in the tracker sheet.",
        "generated": generated, "bank_folder_id": bank_id,
        "tracker": {"name": name, "id": tracker_id},
        "statuses": STATUSES,
        "clips": rows,
    }
    with open(os.path.join(out, "bank_index.json"), "w") as f:
        json.dump(index, f, indent=2)
    with open(os.path.join(out, "README_FOR_AI.md"), "w") as f:
        f.write(readme(rows, tracker_id, bank_id, generated))
    rc("copyto", os.path.join(out, "bank_index.json"), remote(bank_id, "bank_index.json"))
    rc("copyto", os.path.join(out, "README_FOR_AI.md"), remote(bank_id, "README_FOR_AI.md"))
    rc("mkdir", remote(bank_id, "_posted_log"))
    log(f"Tracker: https://docs.google.com/spreadsheets/d/{tracker_id}" if tracker_id else "Tracker uploaded")
